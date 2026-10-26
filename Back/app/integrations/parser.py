import io
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePosixPath
from app.core.errors import BusinessError, require


@dataclass
class ParsedDocument:
    body: str
    images: dict[str, bytes] = field(default_factory=dict)


def parse_document(name, data, mineru_config):
    ext = PurePosixPath(name).suffix.lower()
    if ext in {'.md', '.markdown', '.txt'}:
        for encoding in ['utf-8-sig', 'gb18030']:
            try:
                result = data.decode(encoding)
                require('\x00' not in result, '文本文件包含二进制内容')
                return ParsedDocument(result)
            except UnicodeDecodeError:
                continue
        raise BusinessError(400, '无法识别文本编码，请转换为 UTF-8 后重新上传')
    if ext == '.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                require(len(archive.infolist()) <= 1000 and
                        sum(f.file_size for f in archive.infolist()) <= 100*1024*1024,
                        'Word 解压后的内容超过安全大小限制')
            document = Document(io.BytesIO(data))
            lines = []
            for block in document.iter_inner_content():
                if isinstance(block, Paragraph):
                    style = block.style.name if block.style else ''
                    prefix = ''
                    if style.startswith('Heading ') and style[-1:].isdigit():
                        prefix = '#'*min(6, int(style[-1]))+' '
                    lines.append(prefix+block.text)
                else:
                    for row in block.rows:
                        lines.append('| '+' | '.join(c.text.replace('\n', ' ') for c in row.cells)+' |')
            return ParsedDocument('\n\n'.join(lines))
        except BusinessError:
            raise
        except Exception as exc:
            raise BusinessError(400, 'Word 文件损坏或格式不受支持') from exc
    if ext in {'.pdf', '.doc'}:
        if not mineru_config['enabled']:
            raise BusinessError(503, 'MinerU 尚未启用；PDF/旧版 Word 保留原文件，启用后可在导入中心重试')
        from app.integrations.mineru import MinerUClient
        result = MinerUClient(mineru_config).extract(name, data)
        return ParsedDocument(result.markdown, result.images)
    raise BusinessError(400, '不支持的文档格式')
